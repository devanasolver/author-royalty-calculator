"""Read a KDP royalty report you downloaded and summarize it.

Supports the "Prior Months' Royalties" Excel file (.xlsx) and the
month-to-date CSV from the KDP Reports tab. Nothing connects to your
Amazon account - you download the file, this reads it.
"""

import csv
import os

ROYALTY_WORDS = ("royalty", "earnings", "accrued")
TITLE_WORDS = ("title",)
MARKET_WORDS = ("marketplace", "market")


def _find_columns(header):
    def pick(words):
        lowered = [str(h).strip().lower() for h in header]
        for i, h in enumerate(lowered):      # exact match first
            if h in words:
                return i
        for i, h in enumerate(lowered):      # then "starts with"
            if any(h.startswith(w) for w in words):
                return i
        for i, h in enumerate(lowered):      # then "contains"
            if any(w in h for w in words):
                return i
        return None
    t = pick(TITLE_WORDS)
    r = pick(ROYALTY_WORDS)
    m = pick(MARKET_WORDS)
    if r is not None and any(w in str(header[r]).lower() for w in ("type",)):
        # "Royalty Type" is not the money column - look for a better one
        for i, h in enumerate(header):
            hl = str(h).strip().lower()
            if hl in ROYALTY_WORDS or hl.endswith("royalty"):
                r = i
                break
    return t, r, m


def _num(v):
    if v is None:
        return 0.0
    s = str(v).replace(",", "").replace("$", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def _summarize(rows, tcol, rcol, mcol):
    by_title, by_market, total = {}, {}, 0.0
    for row in rows:
        title = str(row[tcol]).strip() if tcol is not None and tcol < len(row) else "(untitled)"
        market = str(row[mcol]).strip() if mcol is not None and mcol < len(row) else "(all)"
        amt = _num(row[rcol]) if rcol < len(row) else 0.0
        by_title[title] = by_title.get(title, 0) + amt
        by_market[market] = by_market.get(market, 0) + amt
        total += amt
    return {"total": total, "by_title": by_title, "by_market": by_market,
            "rows": len(rows)}


def parse_report(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(
            c is not None and str(c).strip() for c in r)]
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [r for r in csv.reader(f) if any(c.strip() for c in r)]
    else:
        raise ValueError("give me the .xlsx or .csv from KDP Reports")
    if not rows:
        raise ValueError("the file is empty")
    # find the header row (the one containing 'title' and a royalty-ish word)
    hidx = None
    for i, row in enumerate(rows[:20]):
        t, r, _ = _find_columns(row)
        if t is not None and r is not None:
            hidx = i
            break
    if hidx is None:
        raise ValueError("could not find a header row with Title and Royalty "
                         "columns - is this a KDP royalty report?")
    header = rows[hidx]
    t, r, m = _find_columns(header)
    out = _summarize(rows[hidx + 1:], t, r, m)
    out["currency_note"] = ("amounts are in each marketplace's currency - "
                            "KDP pays out converted to your bank's currency")
    return out
